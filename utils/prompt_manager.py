import streamlit as st
import docx
import io
import random
from google_services import (
    get_sheet_service,
    check_sheet_exists,
    create_sheet,
    SPREADSHEET_ID
)

# List of fun random character names
RANDOM_NAMES = [
    "Ziggy", "Nova", "Echo", "Atlas", "Phoenix", "Sage", "Luna", "Orion",
    "Aurora", "Kai", "Zephyr", "Iris", "Thorne", "Raven", "Storm", "Ash"
]

def initialize_system_prompt_state():
    """Initialize system prompt related session state variables"""
    if 'custom_system_prompt' not in st.session_state:
        st.session_state.custom_system_prompt = None
    if 'using_custom_prompt' not in st.session_state:
        st.session_state.using_custom_prompt = False
    if 'character_name' not in st.session_state:
        st.session_state.character_name = "Fred"
    if 'character_emoji' not in st.session_state:
        st.session_state.character_emoji = "👨‍💼"

def read_docx(file):
    """Read text from a .docx file"""
    doc = docx.Document(file)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return '\n'.join(full_text)

def save_character_to_sheet(name, prompt):
    """Save character data to the characters sheet"""
    try:
        sheet_service = get_sheet_service()
        if not sheet_service:
            return False
            
        # Check if characters sheet exists, create if not
        if not check_sheet_exists(sheet_service, SPREADSHEET_ID, "characters"):
            if not create_sheet(sheet_service, SPREADSHEET_ID, "characters"):
                st.error("Failed to create characters sheet")
                return False
            
            # Add headers
            sheet_service.spreadsheets().values().update(
                spreadsheetId=SPREADSHEET_ID,
                range="characters!A1:B1",
                valueInputOption="RAW",
                body={"values": [["Character Name", "System Prompt"]]}
            ).execute()
        
        # Check if character already exists
        result = sheet_service.spreadsheets().values().get(
            spreadsheetId=SPREADSHEET_ID,
            range="characters!A:B"
        ).execute()
        
        values = result.get('values', [])
        row_index = None
        
        # Look for existing character
        for idx, row in enumerate(values):
            if row[0] == name:
                row_index = idx + 1  # +1 because sheets are 1-indexed
                break
        
        if row_index:
            # Update existing character
            range_name = f"characters!A{row_index}:B{row_index}"
            sheet_service.spreadsheets().values().update(
                spreadsheetId=SPREADSHEET_ID,
                range=range_name,
                valueInputOption="RAW",
                body={"values": [[name, prompt]]}
            ).execute()
        else:
            # Add new character
            sheet_service.spreadsheets().values().append(
                spreadsheetId=SPREADSHEET_ID,
                range="characters!A:B",
                valueInputOption="RAW",
                body={"values": [[name, prompt]]}
            ).execute()
        
        return True
    except Exception as e:
        st.error(f"Error saving character: {e}")
        return False

def load_characters():
    """Load all characters from the sheet"""
    try:
        sheet_service = get_sheet_service()
        if not sheet_service:
            return []
            
        if not check_sheet_exists(sheet_service, SPREADSHEET_ID, "characters"):
            return []
            
        result = sheet_service.spreadsheets().values().get(
            spreadsheetId=SPREADSHEET_ID,
            range="characters!A:B"
        ).execute()
        
        values = result.get('values', [])
        if not values:
            return []
            
        # Skip header row and convert to list of tuples
        return [(row[0], row[1]) for row in values[1:] if len(row) >= 2]
    except Exception as e:
        st.error(f"Error loading characters: {e}")
        return []

def get_random_name_and_emoji():
    """Generate a random character name and emoji"""
    name = random.choice(RANDOM_NAMES)
    emojis = ["🤖", "🎭", "🦸", "🧙", "🦊", "🐉", "🎪", "🎯", "🎨", "🎮"]
    emoji = random.choice(emojis)
    return name, emoji

def add_system_prompt_manager():
    """Add system prompt management UI to the sidebar"""
    st.sidebar.markdown("---")
    
    # Simple button in sidebar to trigger configuration
    if st.sidebar.button("⚙️ Custom GPT", use_container_width=True, help="Configure custom GPT settings"):
        st.session_state.show_gpt_config = True
        st.session_state.show_history = False  # Hide history view if open
    
    # Show current character info if using custom
    if st.session_state.using_custom_prompt:
        st.sidebar.caption(f"Using: {st.session_state.character_name} {st.session_state.character_emoji}")

def render_gpt_config():
    """Render the GPT configuration interface in the main page"""
    if not st.session_state.get("show_gpt_config", False):
        return
        
    # Add back button
    col1, col2 = st.columns([0.2, 0.8])
    with col1:
        if st.button("← Back", use_container_width=True):
            st.session_state.show_gpt_config = False
            st.rerun()
    
    st.title("Custom GPT Configuration")
    
    # Radio button to choose between default and custom prompt
    prompt_choice = st.radio(
        "Choose System Prompt:",
        ["Use Default", "Upload Document", "Enter Custom", "Load Saved Character"],
        key="prompt_choice",
        horizontal=True
    )
    
    st.divider()
    
    if prompt_choice == "Use Default":
        st.session_state.using_custom_prompt = False
        st.session_state.custom_system_prompt = None
        st.session_state.character_name = "Fred"
        st.session_state.character_emoji = "👨‍💼"
        
    elif prompt_choice == "Upload Document":
        uploaded_file = st.file_uploader(
            "Upload .docx file",
            type=['docx'],
            key="system_prompt_file"
        )
        
        if uploaded_file:
            try:
                content = read_docx(uploaded_file)
                st.session_state.custom_system_prompt = content
                st.session_state.using_custom_prompt = True
                
                # Character name input
                col1, col2 = st.columns([2, 1])
                with col1:
                    char_name = st.text_input(
                        "Character Name (optional)",
                        key="char_name_upload"
                    )
                
                if not char_name:
                    char_name, emoji = get_random_name_and_emoji()
                    st.info(f"Using random name: {char_name} {emoji}")
                else:
                    emoji = "🎭"
                
                st.session_state.character_name = char_name
                st.session_state.character_emoji = emoji
                
                if st.button("Save Character", use_container_width=True, type="primary"):
                    if save_character_to_sheet(char_name, content):
                        st.success(f"Character {char_name} saved!")
                
                # Show preview in an expander
                with st.expander("Preview Prompt"):
                    st.write(content)
                    
                # Add Load Character button
                if st.button("Load Character", use_container_width=True, type="primary"):
                    st.session_state.custom_system_prompt = content
                    st.session_state.using_custom_prompt = True
                    st.session_state.character_name = char_name
                    st.session_state.character_emoji = emoji
                    st.success(f"Character {char_name} loaded successfully!")
                    st.rerun()
                    
            except Exception as e:
                st.error(f"Error reading file: {str(e)}")
                st.session_state.using_custom_prompt = False
                st.session_state.custom_system_prompt = None
                
    elif prompt_choice == "Enter Custom":
        custom_prompt = st.text_area(
            "Enter custom system prompt",
            height=150,
            key="custom_prompt_input"
        )
        
        if custom_prompt:
            st.session_state.custom_system_prompt = custom_prompt
            st.session_state.using_custom_prompt = True
            
            # Character name input
            col1, col2 = st.columns([2, 1])
            with col1:
                char_name = st.text_input(
                    "Character Name (optional)",
                    key="char_name_custom"
                )
            
            if not char_name:
                char_name, emoji = get_random_name_and_emoji()
                st.info(f"Using random name: {char_name} {emoji}")
            else:
                emoji = "🎭"
            
            st.session_state.character_name = char_name
            st.session_state.character_emoji = emoji
            
            if st.button("Save Character", use_container_width=True, type="primary"):
                if save_character_to_sheet(char_name, custom_prompt):
                    st.success(f"Character {char_name} saved!")
        else:
            st.session_state.using_custom_prompt = False
            st.session_state.custom_system_prompt = None
            
    elif prompt_choice == "Load Saved Character":
        characters = load_characters()
        if characters:
            char_names = [char[0] for char in characters]
            selected_char = st.selectbox(
                "Select Character",
                char_names,
                key="saved_char_select"
            )
            
            # Find the selected character's prompt
            selected_prompt = None
            for name, prompt in characters:
                if name == selected_char:
                    selected_prompt = prompt
                    break
            
            if selected_prompt:
                # Show preview in an expander
                with st.expander("Preview Prompt"):
                    st.write(selected_prompt)
                
                # Add Load Character button outside the expander
                if st.button("Load Character", use_container_width=True, type="primary"):
                    st.session_state.custom_system_prompt = selected_prompt
                    st.session_state.using_custom_prompt = True
                    st.session_state.character_name = selected_char
                    st.session_state.character_emoji = "🎭"
                    st.success(f"Character {selected_char} loaded successfully!")
                    st.rerun()
        else:
            st.info("No saved characters found")

def get_current_system_prompt(default_prompt):
    """Get the current system prompt based on user settings"""
    if st.session_state.using_custom_prompt and st.session_state.custom_system_prompt:
        return st.session_state.custom_system_prompt
    return default_prompt 